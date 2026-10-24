
import pandas as pd
from openpyxl import load_workbook

from universal_merge import merge_files


def create_test_files(folder):
    folder.mkdir()

    df = pd.DataFrame({
        "price": ["12.50", "25.75"],
        "customer_id": ["00123", "00456"]
    })

    df.to_csv(
        folder / "test.csv",
        index=False
    )


def test_keep_original(tmp_path):
    input_folder = tmp_path / "input"
    output_folder = tmp_path / "output"

    create_test_files(input_folder)

    result = merge_files(
        input_folder,
        output_folder
    )

    wb = load_workbook(result["report"])
    ws = wb["Merged Data 1"]

    # Price remains text without user approval
    assert ws["A2"].data_type == "s"

    # Customer ID keeps leading zeros
    assert ws["B2"].value == "00123"


def test_user_selected_number(tmp_path):
    input_folder = tmp_path / "input"
    output_folder = tmp_path / "output"

    create_test_files(input_folder)

    result = merge_files(
        input_folder,
        output_folder,
        column_types={
            "price": "number"
        }
    )

    wb = load_workbook(result["report"])
    ws = wb["Merged Data 1"]

    # Price becomes a real Excel number
    assert ws["A2"].data_type == "n"
    assert ws["A2"].value == 12.50

    # Customer ID remains unchanged
    assert ws["B2"].value == "00123"
    assert ws["B2"].data_type == "s"
